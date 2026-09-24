"""Excel-backed storage for the Lokamania 09 site.

The workbook is the single source of truth. Design rules:

* Every operation loads the workbook fresh from disk (no stale state).
* Every save is atomic: write to a temp file in the same folder, then
  os.replace() over the real file.
* A threading.Lock serializes writes so concurrent submissions can never
  interleave and corrupt the file.
* The server refuses to save while Excel appears to have the file open
  (a ~$lock file exists) so it never silently clobbers Excel's copy.
* A timestamped backup of the workbook is made before the first write
  of each server run.
"""

from __future__ import annotations

import datetime as dt
import os
import re
import shutil
import threading
from pathlib import Path

import openpyxl

APP_SHEET = "Applications"
SETTINGS_SHEET = "Settings"
TABLE_NAME = "LokamaniaApplications"
HEADER_ROW = 6
FIRST_DATA_ROW = 7
MAX_ROWS = 1006  # pre-validated rows; also the range the Dashboard counts
NUM_COLS = 24

# Normalized header (Applications row 6) -> python field key
FIELD_BY_HEADER = {
    "application reference": "reference",
    "submitted at": "submitted_at",
    "brand name": "brand",
    "brand launch year": "launch_year",
    "commercial registration status": "registration",
    "contact person — full name": "contact",
    "phone / whatsapp": "phone",
    "email": "email",
    "instagram": "instagram",
    "website (required)": "website",
    "brand description": "about_brand",
    "main category": "category",
    "subcategory": "subcategory",
    "participation days": "days",
    "requested space": "booth",
    "products": "products",
    "typical price range": "prices",
    "winter collection description": "winter",
    "planned launch / exclusive products": "exclusive",
    "previous participation": "joined",
    "application status": "status",
    "internal notes": "notes",
    "payment status": "payment",
    "last updated": "last_updated",
}

DATETIME_FIELDS = ("submitted_at", "last_updated")


class ExcelWriteRefused(Exception):
    """Raised when we must not write (e.g. Excel has the file open)."""


class ExcelDatabase:
    def __init__(self, path, backup_dir=None):
        self.path = Path(path)
        self.backup_dir = Path(backup_dir or (self.path.parent / "backups"))
        self._lock = threading.Lock()
        self._backup_done = False

    # --------------------------------------------------------------- loading
    def _load(self):
        if not self.path.exists():
            raise FileNotFoundError(f"Workbook not found: {self.path}")
        wb = openpyxl.load_workbook(self.path, data_only=False)
        return wb, wb[APP_SHEET]

    @staticmethod
    def _norm(value):
        return " ".join(str(value).split()).lower()

    def _columns(self, ws):
        """python field key -> 1-based column index, from the header row."""
        cols = {}
        for cell in ws[HEADER_ROW]:
            if cell.value is None:
                continue
            key = FIELD_BY_HEADER.get(self._norm(cell.value))
            if key:
                cols[key] = cell.column
        return cols

    @staticmethod
    def _column_values(ws, col, start=FIRST_DATA_ROW, limit=40):
        """Read a contiguous list of values from a column (stops at first gap)."""
        values = []
        for row in range(start, start + limit):
            v = ws.cell(row, col).value
            if v is None or str(v).strip() == "":
                break
            values.append(str(v).strip())
        return values

    def settings(self):
        """Read the Settings sheet (single source of truth for options)."""
        wb, _ = self._load()
        return self._settings(wb)

    def _settings(self, wb):
        ws = wb[SETTINGS_SHEET]
        categories, spaces_by_category = [], {}
        row = FIRST_DATA_ROW
        while True:
            cat = ws.cell(row, 1).value
            if cat is None or str(cat).strip() == "":
                break
            cat = str(cat).strip()
            categories.append(cat)
            spaces = ws.cell(row, 2).value
            spaces_by_category[cat] = (
                [s.strip() for s in str(spaces).split(",") if s.strip()]
                if spaces
                else []
            )
            row += 1
        days = self._column_values(ws, 4)
        registration = self._column_values(ws, 6)
        joined = self._column_values(ws, 8)
        statuses = self._column_values(ws, 10)
        payments = self._column_values(ws, 12)
        return {
            "categories": categories,
            "spaces_by_category": spaces_by_category,
            "days": days,
            "registration": registration,
            "joined": joined,
            "statuses": statuses,
            "payments": payments,
            "default_status": statuses[0] if statuses else "New",
            "default_payment": payments[0] if payments else "Not Requested",
        }

    # -------------------------------------------------------------- scanning
    @staticmethod
    def _is_filled(ws, row, num_cols=NUM_COLS):
        for col in range(1, num_cols + 1):
            v = ws.cell(row, col).value
            if v is not None and str(v).strip() != "":
                return True
        return False

    @staticmethod
    def _last_data_row(ws):
        last = FIRST_DATA_ROW - 1
        for row in range(FIRST_DATA_ROW, MAX_ROWS + 1):
            if ExcelDatabase._is_filled(ws, row):
                last = row
        return last

    @staticmethod
    def _find_row(ws, reference):
        wanted = str(reference).strip().lower()
        for row in range(FIRST_DATA_ROW, ExcelDatabase._last_data_row(ws) + 1):
            value = ws.cell(row, 1).value
            if value is not None and str(value).strip().lower() == wanted:
                return row
        return None

    @staticmethod
    def _next_reference(ws):
        """Highest existing LM09-XXXX suffix + 1 (e.g. LM09-0004)."""
        highest = 0
        for row in range(FIRST_DATA_ROW, ExcelDatabase._last_data_row(ws) + 1):
            value = ws.cell(row, 1).value
            if not value:
                continue
            match = re.search(r"(\d+)\s*$", str(value).strip())
            if match:
                highest = max(highest, int(match.group(1)))
        return f"LM09-{highest + 1:04d}"

    # ------------------------------------------------------------ write side
    def _ensure_backup(self):
        if self._backup_done:
            return
        stamp = dt.datetime.now().strftime("%Y%m%d_%H%M%S")
        self.backup_dir.mkdir(parents=True, exist_ok=True)
        dest = self.backup_dir / f"{self.path.stem}_backup_{stamp}.xlsx"
        shutil.copy2(self.path, dest)
        self._backup_done = True

    def _check_open_in_excel(self):
        lock = self.path.parent / ("~$" + self.path.name)
        if lock.exists():
            raise ExcelWriteRefused(
                "The workbook appears to be open in Excel. "
                "Close the file and try again."
            )

    def _save(self, wb):
        self._check_open_in_excel()
        tmp = self.path.parent / f".{self.path.stem}.tmp.xlsx"
        try:
            wb.save(tmp)
            os.replace(tmp, self.path)
        finally:
            if tmp.exists():
                try:
                    tmp.unlink()
                except OSError:
                    pass

    @staticmethod
    def _expand_table(ws, last_data_row):
        """Grow the LokamaniaApplications table so new rows are included."""
        table = ws.tables.get(TABLE_NAME)
        if table is None:
            return
        start = table.ref.split(":")[0]
        table.ref = f"{start}:X{last_data_row}"
        if table.autoFilter is not None:
            table.autoFilter.ref = table.ref

    def append_application(self, data):
        """Append one application row and return its Application Reference."""
        with self._lock:
            self._ensure_backup()
            wb, ws = self._load()
            cols = self._columns(ws)
            reference = self._next_reference(ws)
            row = self._last_data_row(ws) + 1
            if row > MAX_ROWS:
                raise ExcelWriteRefused(
                    "The Applications sheet is full (max 1000 rows)."
                )
            now = dt.datetime.now()
            values = dict(data, reference=reference, submitted_at=now,
                          last_updated=now, status=data.get("status"),
                          payment=data.get("payment"))
            style_from = row - 1 if row - 1 >= FIRST_DATA_ROW else None
            for key, col in cols.items():
                if key not in values:
                    continue
                cell = ws.cell(row, col)
                value = values.get(key)
                cell.value = None if value in ("", None) else value
                if style_from is not None:
                    cell._style = ws.cell(style_from, col)._style
                if key in DATETIME_FIELDS:
                    cell.number_format = "yyyy-mm-dd hh:mm"
            self._expand_table(ws, row)
            self._save(wb)
            return reference

    def update_application(self, reference, *, status=None, payment=None,
                           notes=None):
        """Update an existing application row (admin edits)."""
        with self._lock:
            self._ensure_backup()
            wb, ws = self._load()
            cols = self._columns(ws)
            row = self._find_row(ws, reference)
            if row is None:
                raise KeyError(f"No application with reference {reference!r}")
            if status is not None and "status" in cols:
                ws.cell(row, cols["status"]).value = status
            if payment is not None and "payment" in cols:
                ws.cell(row, cols["payment"]).value = payment
            if notes is not None and "notes" in cols:
                ws.cell(row, cols["notes"]).value = notes
            if "last_updated" in cols:
                cell = ws.cell(row, cols["last_updated"])
                cell.value = dt.datetime.now()
                cell.number_format = "yyyy-mm-dd hh:mm"
            self._save(wb)
            return True

    # -------------------------------------------------------------- read side
    def list_applications(self):
        """All application rows, newest first (no lock needed: read-only)."""
        wb, ws = self._load()
        cols = self._columns(ws)
        records = []
        for row in range(FIRST_DATA_ROW, self._last_data_row(ws) + 1):
            record = {}
            for key, col in cols.items():
                value = ws.cell(row, col).value
                if isinstance(value, (dt.datetime, dt.date)):
                    value = value.isoformat(timespec="seconds")
                record[key] = value
            if not (record.get("reference") or record.get("brand")):
                continue
            record["row"] = row
            records.append(record)
        records.sort(key=lambda r: r.get("submitted_at") or "", reverse=True)
        return records

    def stats(self):
        """Counts that mirror the Dashboard sheet (computed from raw rows)."""
        records = self.list_applications()
        settings = self.settings()

        def tally(key, options=None):
            counts = {opt: 0 for opt in (options or [])}
            for record in records:
                value = record.get(key)
                counts[value] = counts.get(value, 0) + 1
            return counts

        by_status = tally("status", settings["statuses"])
        return {
            "total": len(records),
            "by_status": by_status,
            "by_payment": tally("payment", settings["payments"]),
            "by_category": tally("category", settings["categories"]),
            "by_days": tally("days", settings["days"]),
            "by_registration": tally("registration", settings["registration"]),
            "accepted": by_status.get("Accepted", 0),
            "confirmed": by_status.get("Confirmed", 0),
            "payment_pending": by_status.get("Payment Pending", 0),
        }
